import json, os, re
from datetime import datetime, timedelta, timezone
from urllib.parse import quote
from urllib.request import Request, urlopen
from openpyxl import Workbook

BASE="https://www.komtender.ru/api/tenders/get/"
KEY=os.environ.get("KOMTENDER_API_KEY")
TEMPLATE=os.environ.get("KOMTENDER_SEARCH_TEMPLATE_ID","1")
MAX_PAGES=int(os.environ.get("KOMTENDER_PAGES","10"))
if not KEY: raise SystemExit("KOMTENDER_API_KEY is not configured")

# Live runner contract: never print or persist the API key.

def get(path):
    req=Request(BASE+path,headers={"Accept":"application/json","X-API-KEY":KEY})
    with urlopen(req,timeout=45) as r: return json.loads(r.read().decode())

def norm(s):
    return re.sub(r"[\s\-–—_/()\[\],.;:]+","",str(s or "").lower().replace("ё","е"))

def deadline(s):
    if not s: return None
    m=re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:\s+(\d{2}):(\d{2})(?::(\d{2}))?)?",str(s))
    if not m: return None
    return datetime(int(m[1]),int(m[2]),int(m[3]),int(m[4] or 23),int(m[5] or 59),int(m[6] or 59),tzinfo=timezone(timedelta(hours=3)))

materials=json.load(open("data/chzsi_materials.json",encoding="utf-8"))["materials"]
terms=[(m,norm(m)) for m in materials]
now=datetime.now(timezone.utc); cutoff=now+timedelta(hours=72)
info=get("info"); ip=info.get("data",info); remaining=int(ip.get("remaining",0) or 0)
if remaining < 2: raise SystemExit(f"KomTender quota is insufficient: remaining={remaining}")

rows=[]; pages=0
for page in range(1,MAX_PAGES+1):
    if remaining-pages-1<=0: break
    res=get(f"template/{quote(TEMPLATE)}?page={page}&sort=new-first"); pages+=1
    data=res.get("data",[])
    if not isinstance(data,list) or not data: break
    kept=0
    for x in data[:100]:
        d=deadline(x.get("dte"))
        if d is None or d>=cutoff:
            rows.append(x); kept+=1
    if kept==0: break

max_details=max(0,remaining-pages-1)
rows=rows[:max_details]
out_rows=[]
for x in rows:
    try: d=get(quote(str(x.get("id"))))
    except Exception: continue
    cust=d.get("customer") or {}
    positions=d.get("positions") or []
    text=" ".join([str(d.get("descr") or ""),str(cust.get("name") or "")]+[str(p.get("name") or "")+" "+str(p.get("unit") or "") for p in positions])
    nt=norm(text)
    matched=sorted({orig for orig,nm in terms if nm and nm in nt})
    dl=deadline(d.get("dte"))
    if not matched or (dl and dl<cutoff): continue
    qty=[]
    for p in positions:
        if p.get("quantity") is not None: qty.append(f'{p.get("quantity")} {p.get("unit") or ""}'.strip())
    out_rows.append({
      "Тендер":d.get("id"),"Заказчик":cust.get("name",""),"ИНН":cust.get("inn",""),
      "Совпавшие материалы":", ".join(matched),"Описание":d.get("descr",""),
      "Позиции":" | ".join(str(p.get("name") or "") for p in positions),
      "Количество":" | ".join(qty),"НМЦ, ₽":(d.get("price") or {}).get("value"),
      "Регион":d.get("regions",""),"Место":d.get("place",""),"Начало":d.get("dts",""),
      "Окончание (МСК)":d.get("dte"),"Стадия":d.get("stage",""),
      "Ссылка":d.get("url") or f'https://www.komtender.ru/tender/{d.get("id")}',
      "Часов до окончания":round((dl-now).total_seconds()/3600,1) if dl else None})

uniq={str(r["Тендер"]):r for r in out_rows}
out_rows=sorted(uniq.values(),key=lambda r:(r["Окончание (МСК)"] or "9999",str(r["Тендер"])))
os.makedirs("artifacts",exist_ok=True); path="artifacts/chzsi_komtender_material_tenders.xlsx"
wb=Workbook(); ws=wb.active; ws.title="Тендеры ЧЗСИ"
headers=["Тендер","Заказчик","ИНН","Совпавшие материалы","Описание","Позиции","Количество","НМЦ, ₽","Регион","Место","Начало","Окончание (МСК)","Стадия","Ссылка","Часов до окончания"]
ws.append(headers)
for r in out_rows: ws.append([r.get(h) for h in headers])
for c in ws[1]: c.font=c.font.copy(bold=True)
for i,w in enumerate([14,40,16,45,55,65,28,16,28,30,22,24,20,55,20],1): ws.column_dimensions[chr(64+i)].width=w
ws.freeze_panes="A2"
for row in ws.iter_rows(min_row=2):
    if row[13].value: row[13].hyperlink=row[13].value; row[13].style="Hyperlink"
meta=wb.create_sheet("Параметры")
for r in [["Проверка UTC",now.isoformat()],["Порог 72 часа UTC",cutoff.isoformat()],["Шаблон",TEMPLATE],["Страниц",pages],["Карточек",len(rows)],["Совпавших тендеров",len(out_rows)],["Остаток API перед сканированием",remaining],["Источник материалов","Приложенные Excel + PDF ЧЗСИ"],["Фильтр","Марки/сплавы из приложенных материалов; срок подачи >= 72 часов; все заказчики; без веб-поиска"]]: meta.append(r)
meta.column_dimensions["A"].width=34; meta.column_dimensions["B"].width=100
wb.save(path)
print(json.dumps({"ok":True,"matched":len(out_rows),"pages":pages,"details":len(rows),"remaining":remaining,"cutoff":cutoff.isoformat(),"file":path},ensure_ascii=False))
